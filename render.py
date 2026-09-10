#!/usr/bin/env python3
"""
render.py — Rendu des blocs extraits du .docx.

Trois sorties distinctes, depuis le même modèle de blocs :

  * `render_prompt_text`  → le texte envoyé à Claude (marqueurs [IMAGE n] /
    [TABLEAU n] + tableau en markdown : compact en tokens, lisible par le modèle).
  * `render_table_html`   → HTML auto-suffisant (CSS inline) injecté dans la carte
    Anki à l'export, via le placeholder {{TABLE:n}}.
  * `render_preview_html` → aperçu dans l'interface (images servies par /media/...).

`resolve_placeholders` fait la substitution finale {{IMG:n}} / {{TABLE:n}} dans la
réponse de Claude : le modèle ne manipule jamais de nom de fichier, donc il ne
peut pas en inventer.
"""

import html
import re

from docx.text.paragraph import Paragraph

# ──────────────────────────────────────────────────────────────
# Formatage inline (runs → HTML minimal)
# ──────────────────────────────────────────────────────────────
# Surlignages Word → couleurs CSS (cohérent avec le system prompt : plum/yellow)
HIGHLIGHT_CSS = {
    "YELLOW": "yellow",
    "BRIGHT_GREEN": "lightgreen",
    "TURQUOISE": "paleturquoise",
    "PINK": "plum",
    "BLUE": "lightskyblue",
    "RED": "#ff9999",
    "DARK_YELLOW": "#d4c05a",
    "GRAY_25": "#e0e0e0",
    "GRAY_50": "#c0c0c0",
    "TEAL": "#6fbfbf",
    "GREEN": "#8fd48f",
    "VIOLET": "#d0a0e0",
}


def _escape(text: str) -> str:
    return html.escape(text or "", quote=False).replace("\t", " ").replace("\n", "<br>")


def _run_html(run) -> str:
    text = _escape(run.text)
    if not text:
        return ""
    font = run.font
    highlight = None
    try:
        if font.highlight_color is not None:
            highlight = HIGHLIGHT_CSS.get(str(font.highlight_color).split()[0].upper())
    except Exception:
        highlight = None

    if run.bold:
        text = "<b>%s</b>" % text
    if run.italic:
        text = "<i>%s</i>" % text
    if run.underline:
        text = "<u>%s</u>" % text
    try:
        if font.superscript:
            text = "<sup>%s</sup>" % text
        elif font.subscript:
            text = "<sub>%s</sub>" % text
    except Exception:
        pass
    if highlight:
        text = '<span style="background-color: %s;">%s</span>' % (highlight, text)
    return text


def runs_to_html(para: Paragraph) -> str:
    """HTML inline d'un paragraphe (gras / italique / souligné / surlignage)."""
    parts = []
    try:
        items = list(para.iter_inner_content())
    except Exception:
        items = list(para.runs)

    for item in items:
        runs = getattr(item, "runs", None)
        if runs is None:
            parts.append(_run_html(item))
        else:
            # Hyperlink : on garde le texte, pas le lien (inutile dans une carte)
            parts.extend(_run_html(r) for r in runs)
    return "".join(p for p in parts if p).strip()


# ──────────────────────────────────────────────────────────────
# Tableaux
# ──────────────────────────────────────────────────────────────
TABLE_STYLE = (
    "border-collapse: collapse; margin: 8px 0; font-size: 0.9em; "
    "text-align: left; max-width: 100%;"
)
CELL_STYLE = "border: 1px solid #999; padding: 4px 8px; vertical-align: top;"
HEAD_STYLE = CELL_STYLE + " background-color: #f0f0f3; font-weight: bold;"


def render_table_html(table: dict, media_prefix: str = "") -> str:
    """
    Tableau → HTML auto-suffisant (CSS inline, pas de dépendance au style Anki).
    `media_prefix` permet de préfixer les images pour l'aperçu web ; à l'export
    Anki il reste vide (collection.media est un espace plat).
    """
    rows_html = []
    for row in table.get("rows", []):
        cells_html = []
        for cell in row:
            tag = "th" if cell.get("header") else "td"
            style = HEAD_STYLE if cell.get("header") else CELL_STYLE
            attrs = ""
            if cell.get("colspan", 1) > 1:
                attrs += ' colspan="%d"' % cell["colspan"]
            if cell.get("rowspan", 1) > 1:
                attrs += ' rowspan="%d"' % cell["rowspan"]
            content = cell.get("html") or _escape(
                CELLIMG_RE.sub("", cell.get("text", "")).strip()
            )
            for nested in cell.get("tables", []):
                content += render_table_html(nested, media_prefix)
            if media_prefix:
                content = _prefix_img_src(content, media_prefix)
            cells_html.append(
                '<%s style="%s"%s>%s</%s>' % (tag, style, attrs, content or "&nbsp;", tag)
            )
        rows_html.append("<tr>%s</tr>" % "".join(cells_html))

    table_html = '<table style="%s">%s</table>' % (TABLE_STYLE, "".join(rows_html))
    # Un tableau large doit pouvoir défiler sur AnkiDroid plutôt que déborder
    return '<div style="overflow-x: auto;">%s</div>' % table_html


def _prefix_img_src(html_fragment: str, prefix: str) -> str:
    return re.sub(
        r'(<img[^>]*\bsrc=")([^"/][^"]*)(")',
        lambda m: m.group(1) + prefix.rstrip("/") + "/" + m.group(2) + m.group(3),
        html_fragment,
    )


CELLIMG_RE = re.compile(r"\{\{CELLIMG:([A-Za-z0-9_]+)\}\}")


def _resolve_cell_images(text: str, assets: dict) -> str:
    """{{CELLIMG:img_3}} → [image: description] (résolu à chaque rendu du prompt)."""
    def _sub(match):
        asset = (assets or {}).get(match.group(1)) or {}
        return "[image: %s]" % (asset.get("alt") or "sans description")

    return CELLIMG_RE.sub(_sub, text or "")


def render_table_text(table: dict, assets: dict = None) -> str:
    """Tableau → markdown compact (pour le prompt : ~4× moins de tokens que le HTML)."""
    rows = table.get("rows", [])
    n_cols = max(1, table.get("n_cols") or 1)
    matrix = []
    header_row_idx = None

    for ri, row in enumerate(rows):
        line = [""] * n_cols
        for cell in row:
            col = cell.get("col", 0)
            if col >= n_cols:
                continue
            text = _resolve_cell_images(cell.get("text") or "", assets)
            text = text.replace("\n", " / ").replace("|", "\\|")
            for nested in cell.get("tables", []):
                text += " [tableau imbriqué : %s]" % (
                    render_table_text(nested, assets).replace("\n", " ; ")
                )
            line[col] = text.strip()
            if cell.get("header") and header_row_idx is None:
                header_row_idx = ri
        matrix.append(line)

    out = []
    for ri, line in enumerate(matrix):
        out.append("| " + " | ".join(line) + " |")
        if ri == (header_row_idx if header_row_idx is not None else -1):
            out.append("|" + "|".join(["---"] * n_cols) + "|")
    return "\n".join(out)


# ──────────────────────────────────────────────────────────────
# Prompt envoyé à Claude
# ──────────────────────────────────────────────────────────────
def render_prompt_text(blocks: list, assets: dict) -> str:
    """
    Rend les blocs d'un chunk en texte destiné à Claude.
    Les images deviennent des marqueurs décrits par leur texte alternatif ;
    les tableaux sont inclus en markdown avec leur placeholder de réutilisation.
    """
    parts = []
    for block in blocks:
        btype = block.get("type")

        if btype == "text":
            parts.append(block.get("text", ""))

        elif btype == "image":
            n = block.get("n", 1)
            asset = assets.get(block.get("asset"), {})
            alt = (asset.get("alt") or "").strip()
            if asset.get("missing"):
                continue  # image illisible : inutile d'en parler au modèle
            if alt:
                parts.append(
                    "[IMAGE %d — description : %s]\n"
                    "(écris {{IMG:%d}} dans une carte pour y afficher cette image)" % (n, alt, n)
                )
            else:
                parts.append(
                    "[IMAGE %d — aucune description disponible]\n"
                    "(écris {{IMG:%d}} dans une carte pour y afficher cette image)" % (n, n)
                )

        elif btype == "table":
            n = block.get("n", 1)
            parts.append(
                "[TABLEAU %d]\n%s\n"
                "(écris {{TABLE:%d}} pour réutiliser ce tableau dans une carte)"
                % (n, render_table_text(block, assets), n)
            )

    return "\n\n".join(p for p in parts if p and p.strip()).strip()


# ──────────────────────────────────────────────────────────────
# Aperçu dans l'interface
# ──────────────────────────────────────────────────────────────
def render_preview_html(blocks: list, assets: dict, media_prefix: str) -> str:
    """Aperçu d'un chunk dans l'app (texte + vignettes + tableaux rendus)."""
    out = []
    for block in blocks:
        btype = block.get("type")

        if btype == "text":
            out.append('<p class="pv-text">%s</p>' % _escape(block.get("text", "")))

        elif btype == "image":
            n = block.get("n", 1)
            asset = assets.get(block.get("asset")) or {}
            if asset.get("missing") or not asset.get("file"):
                out.append(
                    '<div class="pv-image pv-image--missing" data-asset="%s">'
                    "<span>IMAGE %d — %s</span></div>"
                    % (html.escape(asset.get("id", "")), n,
                       html.escape(asset.get("note") or "image illisible"))
                )
                continue
            src = "%s/%s" % (media_prefix.rstrip("/"), asset["file"])
            alt = asset.get("alt") or ""
            out.append(
                '<figure class="pv-image" data-asset="%s" data-n="%d">'
                '<img src="%s" alt="%s" loading="lazy">'
                '<figcaption>{{IMG:%d}} · %s</figcaption>'
                "</figure>"
                % (
                    html.escape(asset.get("id", "")),
                    n,
                    html.escape(src, quote=True),
                    html.escape(alt, quote=True),
                    n,
                    html.escape(alt) if alt else "<em>aucune description</em>",
                )
            )

        elif btype == "table":
            n = block.get("n", 1)
            out.append(
                '<div class="pv-table" data-n="%d"><div class="pv-table-tag">{{TABLE:%d}}</div>%s</div>'
                % (n, n, render_table_html(block, media_prefix))
            )

    return "\n".join(out)


# ──────────────────────────────────────────────────────────────
# Substitution des placeholders
# ──────────────────────────────────────────────────────────────
PLACEHOLDER_RE = re.compile(
    r"\{\{\s*(IMG|IMAGE|TABLE|TABLEAU)\s*[:\-# ]?\s*(\d+)\s*\}\}", re.IGNORECASE
)


def image_html(asset: dict, media_prefix: str = "") -> str:
    if not asset or not asset.get("file"):
        return ""
    src = asset["file"]
    if media_prefix:
        src = "%s/%s" % (media_prefix.rstrip("/"), src)
    alt = html.escape(asset.get("alt") or "", quote=True)
    return '<img src="%s" alt="%s">' % (html.escape(src, quote=True), alt)


def chunk_placeholder_map(chunk: dict) -> dict:
    """{('IMG', 1): block, ('TABLE', 2): block} pour un chunk."""
    mapping = {}
    for block in chunk.get("blocks", []):
        if block.get("type") == "image":
            mapping[("IMG", block.get("n", 1))] = block
        elif block.get("type") == "table":
            mapping[("TABLE", block.get("n", 1))] = block
    return mapping


def resolve_placeholders(text: str, chunk: dict, assets: dict, media_prefix: str = ""):
    """
    Remplace {{IMG:n}} / {{TABLE:n}} dans la réponse de Claude.

    Retourne (texte_résolu, assets_utilisés:set, placeholders_inconnus:list).
    Un placeholder inconnu (hallucination du modèle) est supprimé et signalé.
    """
    mapping = chunk_placeholder_map(chunk)
    used: set = set()
    unknown: list = []

    def _sub(match):
        kind = match.group(1).upper()
        kind = "IMG" if kind in ("IMG", "IMAGE") else "TABLE"
        n = int(match.group(2))
        block = mapping.get((kind, n))
        if block is None:
            unknown.append("{{%s:%d}}" % (kind, n))
            return ""
        if kind == "IMG":
            asset = assets.get(block.get("asset")) or {}
            if not asset.get("file"):
                unknown.append("{{IMG:%d}} (image illisible)" % n)
                return ""
            used.add(asset["id"])
            return image_html(asset, media_prefix)
        for aid in block.get("assets", []):
            used.add(aid)
        return render_table_html(block, media_prefix)

    return PLACEHOLDER_RE.sub(_sub, text or ""), used, unknown
