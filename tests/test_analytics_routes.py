"""Exercise real local workflows and inspect their persisted analytics events."""

import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import app as web
import paths
import project_store
import settings
from telemetry import Telemetry


class AnalyticsRouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        for name in ("uploads", "projects", "media", "exports"):
            (self.root / name).mkdir()
        for module, name, value in (
            (settings, "SETTINGS_FILE", self.root / "settings.json"),
            (project_store, "PROJECTS_DIR", self.root / "projects"),
            (project_store, "MEDIA_DIR", self.root / "media"),
            (paths, "MEDIA_DIR", self.root / "media"),
            (web, "UPLOAD_DIR", self.root / "uploads"),
            (web, "MEDIA_DIR", self.root / "media"),
        ):
            guard = patch.object(module, name, value)
            guard.start()
            self.addCleanup(guard.stop)
        self.analytics = Telemetry(self.root / "analytics-state.json", settings.load_settings,
                                   lambda: {"project_token": "phc_test", "region": "eu", "app_version": "1.2.1"})
        for guard in (patch.object(self.analytics, "_start_worker"), patch.object(web, "analytics", self.analytics, create=True)):
            guard.start()
            self.addCleanup(guard.stop)
        settings.save_settings({"download_dir": str(self.root / "exports")})
        self.client = web.app.test_client()

    def events(self, name):
        if not self.analytics.state_file.exists():
            return []
        data = json.loads(self.analytics.state_file.read_text(encoding="utf-8"))
        return [e for e in data["queue"] if e["event"] == name]

    def create_project(self):
        doc = Document()
        doc.add_heading("Secret heading", 1)
        doc.add_paragraph("Private learning material")
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        response = self.client.post("/upload", data={"docx": (buffer, "private-course.docx")})
        self.assertEqual(response.status_code, 302)
        return response.headers["Location"].split("/")[-1]

    def test_real_save_and_export_events_have_correct_counts_without_content(self):
        self.client.post("/api/settings", json={"analytics_enabled": True})
        project_id = self.create_project()
        project = project_store.load_project(project_id)
        chunk_id = project["prompts"][0]["id"]
        payload = {"id": chunk_id, "response": "Secret question\tSecret answer\nAnother\tAnswer"}
        self.assertEqual(self.client.post(f"/api/save/{project_id}", json=payload).status_code, 200)
        self.client.post(f"/api/save/{project_id}", json=payload)
        for suffix in ("tsv", "zip"):
            response = self.client.get(f"/export.{suffix}/{project_id}")
            self.assertEqual(response.status_code, 200)
            response.close()
        self.assertEqual(len(self.events("project_created")), 1)
        saved = self.events("cards_saved")
        self.assertEqual(len(saved), 1)
        self.assertEqual(saved[0]["properties"]["card_count"], 2)
        exports = self.events("export_completed")
        self.assertEqual([e["properties"]["format"] for e in exports], ["tsv", "zip"])
        self.assertEqual([e["properties"]["card_count"] for e in exports], [2, 2])
        raw = self.analytics.state_file.read_text(encoding="utf-8")
        for private in ("private-course", "Secret", "Private learning", "Another"):
            self.assertNotIn(private, raw)

    def test_settings_enable_and_withdraw_consent(self):
        self.client.get("/")
        self.assertFalse(self.analytics.state_file.exists())
        self.client.post("/api/settings", json={"analytics_enabled": True})
        self.assertEqual(len(self.events("first_launch")), 1)
        self.client.post("/api/settings", json={"analytics_enabled": False})
        self.assertEqual(self.events("first_launch"), [])
        self.create_project()
        self.assertEqual(self.events("project_created"), [])

    def test_successful_pages_record_activity_but_background_checks_do_not(self):
        settings.save_settings({"analytics_enabled": True})
        self.client.get("/api/version")
        self.assertFalse(self.analytics.state_file.exists())
        self.client.get("/")
        self.client.get("/settings")
        self.assertEqual(len(self.events("app_opened")), 2)

    def test_failed_actions_do_not_count_as_success(self):
        settings.save_settings({"analytics_enabled": True})
        self.client.post("/upload", data={"docx": (io.BytesIO(b"invalid"), "invalid.docx")})
        self.client.post("/api/save/missing", json={"id": 1, "response": "Q\tA"})
        self.client.get("/export.tsv/missing")
        self.assertEqual(self.events("project_created"), [])
        self.assertEqual(self.events("cards_saved"), [])
        self.assertEqual(self.events("export_completed"), [])

    def test_broken_telemetry_storage_cannot_break_project_save_or_export(self):
        settings.save_settings({"analytics_enabled": True})
        self.analytics.state_file.mkdir()
        project_id = self.create_project()
        chunk_id = project_store.load_project(project_id)["prompts"][0]["id"]
        self.assertEqual(self.client.post(f"/api/save/{project_id}", json={"id": chunk_id, "response": "Q\tA"}).status_code, 200)
        response = self.client.get(f"/export.tsv/{project_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Q\tA", response.data)
        response.close()
